#!/usr/bin/env python3
"""Cache supplied materials and audit target/option vocabulary with page evidence."""
import hashlib
import json
import re
from pathlib import Path
from docx import Document
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'memory_palace/sources'
OUT.mkdir(parents=True, exist_ok=True)
inputs = ['unit 13.docx', '期中考共同命題Check_list.docx.pdf', '高中英文參考詞彙表(111學年度起適用).pdf']
manifest = {name: hashlib.sha256((ROOT/name).read_bytes()).hexdigest() for name in inputs}
(OUT/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
doc = Document(ROOT/inputs[0])
text = '\n'.join(p.text for p in doc.paragraphs)
text += '\n' + '\n'.join(' | '.join(c.text for c in row.cells) for t in doc.tables for row in t.rows)
(OUT/'unit13.txt').write_text(text+'\n')
check = PdfReader(ROOT/inputs[1])
(OUT/'checklist.txt').write_text('\n'.join(f'PDF PAGE {i+1}\n'+p.extract_text() for i,p in enumerate(check.pages)))
pdf = PdfReader(ROOT/inputs[2])
# Alphabetical part, file pages 65-115. Keep page references, not a guessed CEFR mapping.
pages = [(i+1,p.extract_text()) for i,p in enumerate(pdf.pages) if 64<=i<=114]
exam = json.loads((ROOT/'exams/g11_midterm_v1.json').read_text())
unit = 'wipe twist drip budget client campus cast squeeze scream flame passion presence adventure hint bore bitter clue rush eager spin'.split()
words = set(unit)
for q in exam['items']:
    if q['section']=='vocab': words.update(q['options'])
words.update(exam['passages']['bank']['word_bank'].values())
lemmas = {'flames':'flame','clouds':'cloud','shadows':'shadow','sparks':'spark','unable':'able'}
audit = {}
for word in sorted(words):
    lemma=lemmas.get(word,word)
    pattern=re.compile(r'(?<![A-Za-z])'+re.escape(lemma)+r'\s+([a-z./()]+)\s+([1-6])\b')
    matches=[]
    for page,text in pages:
        for m in pattern.finditer(text):
            matches.append({'page':page,'pos':m.group(1),'level':int(m.group(2)),'entry':m.group(0).replace('\n',' ')})
    audit[word]={'lemma':lemma,'unit13':lemma in unit,'matches':matches}
    if word=='unable':
        audit[word]['note']='unable 是 un- 加 able 的否定衍生詞；參考詞彙表 PDF 第７頁原則不另列此類詞。此處級別是 able 的參照級別，並非 unable 的獨立官方級別。'
(OUT/'vocabulary_audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2)+'\n')
missing=[w for w,v in audit.items() if not v['matches']]
print(f'Indexed {len(inputs)} sources; checked {len(audit)} vocabulary forms.')
print('Unresolved:',missing)
if missing: raise SystemExit(1)
print('Vocabulary choice levels:')
for q in exam['items']:
    if q['section']=='vocab':
        print(q['n'],[(w,audit[w]['matches'][0]['level']) for w in q['options']])
